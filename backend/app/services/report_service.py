import io
from datetime import date
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.units import inch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def generate_pdf_report(report_type, params):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=inch*0.75, leftMargin=inch*0.75,
                            topMargin=inch*0.75, bottomMargin=inch*0.75)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=18, spaceAfter=12,
                                  textColor=colors.HexColor("#1a237e"))
    story.append(Paragraph("CRPF Welfare Monitoring System", title_style))
    story.append(Paragraph(f"Report Type: {report_type.replace('_',' ').title()}", styles["Heading2"]))
    story.append(Paragraph(f"Generated: {date.today()}", styles["Normal"]))
    story.append(Spacer(1, 0.25*inch))

    from ..models.risk_score import RiskScore
    from ..models.personnel import Personnel
    from sqlalchemy import func
    from ..extensions import db

    subq = (RiskScore.query
        .with_entities(RiskScore.personnel_id, func.max(RiskScore.score_date).label("max_date"))
        .group_by(RiskScore.personnel_id).subquery())
    from sqlalchemy.orm import aliased
    scores = (RiskScore.query.join(subq,
        (RiskScore.personnel_id == subq.c.personnel_id) &
        (RiskScore.score_date == subq.c.max_date)).all())

    data = [["Rank", "Unit", "Risk Level", "Stress Score", "Burnout Score", "Anomaly"]]
    for s in scores:
        p = Personnel.query.get(s.personnel_id)
        if p:
            data.append([p.rank or "-", p.unit_name or "-", s.overall_risk,
                         f"{s.stress_score:.1f}", f"{s.burnout_score:.1f}", "Yes" if s.is_anomaly else "No"])

    table = Table(data, colWidths=[1*inch, 1.5*inch, 1.2*inch, 1.2*inch, 1.2*inch, 0.8*inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1a237e")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
    ]))
    story.append(table)
    doc.build(story)
    return buf, f"welfare_report_{date.today()}.pdf"

def generate_excel_report(report_type, params):
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Welfare Report"
    headers = ["Personnel ID", "Rank", "Unit", "Risk Level", "Stress Score", "Burnout Score", "Score Date"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a237e", end_color="1a237e", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    from ..models.risk_score import RiskScore
    from ..models.personnel import Personnel
    from sqlalchemy import func
    subq = (RiskScore.query
        .with_entities(RiskScore.personnel_id, func.max(RiskScore.score_date).label("max_date"))
        .group_by(RiskScore.personnel_id).subquery())
    scores = (RiskScore.query.join(subq,
        (RiskScore.personnel_id == subq.c.personnel_id) &
        (RiskScore.score_date == subq.c.max_date)).all())
    for row_idx, s in enumerate(scores, 2):
        p = Personnel.query.get(s.personnel_id)
        if p:
            ws.append([p.anon_id, p.rank, p.unit_name, s.overall_risk,
                       s.stress_score, s.burnout_score, str(s.score_date)])
    wb.save(buf)
    return buf, f"welfare_report_{date.today()}.xlsx"
